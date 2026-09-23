import io

import docx
import pytest

from app.ingestion.extract import (
    ExtractionError,
    FileKind,
    detect_kind,
    extract,
    looks_like_valid_file,
    readable_heading,
    title_from_filename,
)


def test_detect_kind_by_extension() -> None:
    assert detect_kind("Policy.PDF") is FileKind.PDF
    assert detect_kind("notes.markdown") is FileKind.MARKDOWN
    assert detect_kind("archive.zip") is None


def test_file_signatures_are_checked() -> None:
    assert looks_like_valid_file(FileKind.PDF, b"%PDF-1.7 ...")
    assert not looks_like_valid_file(FileKind.PDF, b"not a pdf")
    assert not looks_like_valid_file(FileKind.TEXT, b"\xff\xfe\xfa")


def test_titles_are_readable() -> None:
    assert title_from_filename("annual_leave-policy.md") == "Annual Leave Policy"
    assert readable_heading("IT SECURITY AND DATA PROTECTION") == "IT Security and Data Protection"
    assert readable_heading("Already Mixed Case") == "Already Mixed Case"


def test_markdown_headings_and_lists() -> None:
    text = (
        "# Travel Policy\n\n## Hotels\n\nUp to **200** per night.\n\n- London: 275\n- Paris: 250\n"
    )
    document = extract(FileKind.MARKDOWN, text.encode())
    assert document.title == "Travel Policy"
    assert [(block.text, block.heading_level) for block in document.blocks] == [
        ("Travel Policy", 1),
        ("Hotels", 2),
        ("Up to 200 per night.", None),
        ("- London: 275", None),
        ("- Paris: 250", None),
    ]


def test_plain_text_heading_detection() -> None:
    text = (
        "ANNUAL LEAVE POLICY\n\n1. ENTITLEMENT\nAll full-time employees are entitled to 25 days.\n"
        "They also receive public holidays.\n\n2. CARRY-OVER RULES\nUp to 5 days carry over.\n"
    )
    blocks = extract(FileKind.TEXT, text.encode()).blocks
    assert [block.heading_level for block in blocks] == [1, 2, None, 2, None]
    assert blocks[2].text.endswith("They also receive public holidays.")


def test_docx_styles_become_headings() -> None:
    source = docx.Document()
    source.add_heading("Expenses", level=0)
    source.add_heading("Meals", level=2)
    source.add_paragraph("Meals are reimbursed up to 60 per day.")
    buffer = io.BytesIO()
    source.save(buffer)

    document = extract(FileKind.DOCX, buffer.getvalue())
    assert document.title == "Expenses"
    assert document.blocks[1].heading_level == 2
    assert document.blocks[2].text == "Meals are reimbursed up to 60 per day."


def test_corrupt_pdf_is_reported() -> None:
    with pytest.raises(ExtractionError):
        extract(FileKind.PDF, b"%PDF-1.4 truncated")


def test_document_without_body_text_is_rejected() -> None:
    with pytest.raises(ExtractionError):
        extract(FileKind.MARKDOWN, b"# Only a title\n")
