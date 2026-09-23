import io
import re
from dataclasses import dataclass
from enum import StrEnum
from pathlib import PurePath

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class FileKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    MARKDOWN = "md"
    TEXT = "txt"


CONTENT_TYPES = {
    FileKind.PDF: "application/pdf",
    FileKind.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    FileKind.MARKDOWN: "text/markdown",
    FileKind.TEXT: "text/plain",
}


class ExtractionError(Exception):
    pass


@dataclass(frozen=True)
class Block:
    text: str
    heading_level: int | None = None
    page: int | None = None

    @property
    def is_heading(self) -> bool:
        return self.heading_level is not None


@dataclass(frozen=True)
class ExtractedDocument:
    blocks: list[Block]
    page_count: int | None = None

    @property
    def title(self) -> str | None:
        if self.blocks and self.blocks[0].heading_level == 1:
            return readable_heading(self.blocks[0].text)
        return None


MARKDOWN_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")
NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+([A-Z][^.!?]{1,80})$")
LIST_ITEM = re.compile(r"^\s*(?:[-*•]|\d+[.)])\s+")


def detect_kind(filename: str) -> FileKind | None:
    suffix = PurePath(filename).suffix.lower().lstrip(".")
    aliases = {"markdown": "md", "text": "txt"}
    try:
        return FileKind(aliases.get(suffix, suffix))
    except ValueError:
        return None


def looks_like_valid_file(kind: FileKind, data: bytes) -> bool:
    if kind is FileKind.PDF:
        return data.startswith(b"%PDF")
    if kind is FileKind.DOCX:
        return data.startswith(b"PK\x03\x04")
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return True


MINOR_WORDS = {"a", "an", "and", "for", "in", "of", "on", "or", "the", "to"}


def readable_heading(text: str) -> str:
    if not text.isupper():
        return text
    words = []
    for index, word in enumerate(text.split()):
        lowered = word.lower()
        if index and lowered in MINOR_WORDS:
            words.append(lowered)
        elif len(word) <= 3 and word.isalpha():
            words.append(word)
        else:
            words.append(word.capitalize())
    return " ".join(words)


def title_from_filename(filename: str) -> str:
    stem = PurePath(filename).stem
    words = re.sub(r"[_\-]+", " ", stem).split()
    return " ".join(word if word.isupper() else word.capitalize() for word in words) or filename


def plain_heading_level(line: str) -> int | None:
    stripped = line.strip()
    if not stripped or len(stripped) > 90 or stripped.endswith((".", ",", ";", ":")):
        return None
    if NUMBERED_HEADING.match(stripped):
        return 2
    letters = [char for char in stripped if char.isalpha()]
    if len(letters) >= 3 and all(char.isupper() for char in letters):
        return 1 if len(stripped.split()) <= 6 else 2
    return None


def split_plain_text(text: str, page: int | None = None) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block(" ".join(paragraph), page=page))
            paragraph.clear()

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            flush()
            continue
        level = plain_heading_level(line)
        if level is not None:
            flush()
            blocks.append(Block(line, heading_level=level, page=page))
        elif LIST_ITEM.match(line):
            flush()
            blocks.append(Block(line, page=page))
        else:
            paragraph.append(line)
    flush()
    return blocks


def split_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.extend(split_plain_text("\n".join(buffer)))
            buffer.clear()

    for line in text.splitlines():
        match = MARKDOWN_HEADING.match(line.strip())
        if match:
            flush()
            blocks.append(Block(match.group(2).strip(), heading_level=len(match.group(1))))
        else:
            buffer.append(line)
    flush()
    return [
        Block(re.sub(r"\*\*|__|`", "", block.text), block.heading_level, block.page)
        for block in blocks
    ]


def extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError) as exc:
        raise ExtractionError("The PDF could not be read") from exc
    blocks = [
        block for number, text in enumerate(pages, 1) for block in split_plain_text(text, number)
    ]
    return ExtractedDocument(blocks=blocks, page_count=len(pages))


def extract_docx(data: bytes) -> ExtractedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("The Word document could not be read") from exc
    blocks: list[Block] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style == "Title":
            blocks.append(Block(text, heading_level=1))
        elif style.startswith("Heading"):
            digits = "".join(char for char in style if char.isdigit())
            blocks.append(Block(text, heading_level=int(digits) if digits else 2))
        else:
            blocks.append(Block(text))
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(Block(" | ".join(cells)))
    return ExtractedDocument(blocks=blocks)


def extract(kind: FileKind, data: bytes) -> ExtractedDocument:
    if kind is FileKind.PDF:
        extracted = extract_pdf(data)
    elif kind is FileKind.DOCX:
        extracted = extract_docx(data)
    else:
        text = data.decode("utf-8-sig")
        blocks = split_markdown(text) if kind is FileKind.MARKDOWN else split_plain_text(text)
        extracted = ExtractedDocument(blocks=blocks)
    if not any(not block.is_heading for block in extracted.blocks):
        raise ExtractionError("No readable text was found in the document")
    return extracted
